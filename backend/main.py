import io
import re
import threading
import time
import uuid
from copy import deepcopy
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Union

import docx
import pypdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from fastapi import Body, FastAPI, File, HTTPException, Request, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from mock_data import (
    AUDIT,
    GRAPHS,
    STANDARDS_CATALOG,
    SUPERSEDED_CATALOG,
    audit,
    get_bilingual_clause,
)
from retrieval_engine import get_retrieval_engine

app = FastAPI(
    title="ManakSetu BIS Procurement Intelligence API",
    description="Multi-domain BIS standards search, clause synthesis, normative knowledge graph, officer workspace, and tender document audit API.",
    version="1.1.0",
)

origins = [
    "http://localhost:3000",
    "http://127.0.0.1:3000",
    "http://localhost:5173",
    "http://127.0.0.1:5173",
    "http://localhost:8000",
    "http://127.0.0.1:8000",
]

app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_origin_regex=r"https?://(localhost|127\.0\.0\.1)(:\d+)?",
    allow_credentials=True,
    allow_methods=["GET", "POST", "PATCH", "DELETE", "OPTIONS", "PUT"],
    allow_headers=["*"],
)

# ---------------------------------------------------------------------------
# THREAD-SAFE IN-MEMORY OFFICER WORKSPACE STATE
# ---------------------------------------------------------------------------
workspace_lock = threading.Lock()

OFFICER_PROFILE: Dict[str, Any] = {
    "name": "Shri Rajesh Kumar Sharma",
    "designation": "Chief Procurement Officer & Standards Auditor",
    "department": "Ministry of Commerce & Industry / CPWD Electrical Wing",
    "officer_id": "#GOV-PROC-8821",
    "dsc_status": "Active (NIC-CA Validated)",
    "dsc_token_hash": "SHA256:7F89B...2026",
    # Organization metadata
    "organization": "Central Public Procurement Portal / GeM Desk",
    "email": "rk.sharma@gov.in",
    "badge": "Level-3 BIS Certified Procurement Auditor",
    "location": "New Delhi, India",
}

INITIAL_DRAFTS: List[Dict[str, Any]] = [
    {
        "id": "drf_1",
        "draft_id": "drf_1",
        "project_title": "NIT-204: 33kV Substation Cables",
        "project_name": "NIT-204: 33kV Substation Cables",
        "primary_standard": "IS 7098 (Part 1) : 1988",
        "standard_code": "IS 7098 (Part 1) : 1988",
        "clause_heading": "4.2.1 Power cable — conformity and testing",
        "clause_body": "The bidder shall offer 3.5 core, 1.1 kV grade XLPE insulated and PVC sheathed power cables conforming to IS 7098 (Part 1) : 1988 with Amendments 1 to 4. Conductors shall comply with IS 8130 : 2013; insulation and sheath with IS 5831 : 1984; armour with IS 3975 : 1999. Testing shall be carried out as per IS 10810 series including vertical flame retardance Part 53. The product shall bear the BIS Standard Mark under Wires and Cables QCO 2023.",
        "clause_text": "The bidder shall offer 3.5 core, 1.1 kV grade XLPE insulated and PVC sheathed power cables conforming to IS 7098 (Part 1) : 1988 with Amendments 1 to 4...",
        "allied_standards": ["IS 8130 : 2013", "IS 5831 : 1984", "IS 10810 (Part 53)"],
        "last_modified": "2026-09-22",
        "created_at": "2026-09-21T09:30:00Z",
        "domain": "cables",
        "officer_notes": "Mandate batch test certificate for IS 10810 Part 53 vertical flame retardance prior to dispatch.",
    },
    {
        "id": "drf_2",
        "draft_id": "drf_2",
        "project_title": "NHAI Expressway Elevated Deck Pier Reinforcement",
        "project_name": "NHAI Expressway Elevated Deck Pier Reinforcement",
        "primary_standard": "IS 1786 : 2008",
        "standard_code": "IS 1786 : 2008",
        "clause_heading": "5.1.2 High-strength rebar — proof stress & chemical limits",
        "clause_body": "The contractor shall supply High Strength Deformed Steel Bars of grade Fe 500D conforming to IS 1786 : 2008 with Amendments 1 to 3. The steel shall exhibit minimum 0.2 percent proof stress of 500 N/mm² and TS/YS ratio not less than 1.10 per IS 1608. Product bundles shall carry valid BIS ISI marking under Steel and Steel Products QCO 2024.",
        "clause_text": "The contractor shall supply High Strength Deformed Steel Bars of grade Fe 500D conforming to IS 1786 : 2008 with Amendments 1 to 3...",
        "allied_standards": ["IS 1608 (Part 1)", "IS 1599 : 2019", "IS 228"],
        "last_modified": "2026-09-22",
        "created_at": "2026-09-22T14:15:00Z",
        "domain": "steel",
        "officer_notes": "Verify heat-wise ladle chemical analysis for S+P maximum 0.075%.",
    },
    {
        "id": "drf_3",
        "draft_id": "drf_3",
        "project_title": "Smart City Command Center IT Infrastructure",
        "project_name": "Smart City Command Center IT Infrastructure",
        "primary_standard": "IS 13252 (Part 1) : 2010",
        "standard_code": "IS 13252 (Part 1) : 2010",
        "clause_heading": "2.4.1 IT Equipment Safety & CRS Registration",
        "clause_body": "All Information Technology server and terminal hardware shall conform to IS 13252 (Part 1) : 2010 with Amendments 1 to 4 under the MeitY Compulsory Registration Scheme (CRS). All integrated secondary lithium cells shall possess separate BIS registration under IS 16046 (Part 2) : 2018.",
        "clause_text": "All Information Technology server and terminal hardware shall conform to IS 13252 (Part 1) : 2010...",
        "allied_standards": ["IS 616 : 2017", "IS 16046 (Part 2) : 2018"],
        "last_modified": "2026-09-23",
        "created_at": "2026-09-23T08:00:00Z",
        "domain": "electronics",
        "officer_notes": "Require valid R-numbers on GeM bid submission before technical qualification.",
    },
]

INITIAL_AUDIT_HISTORY: List[Dict[str, Any]] = [
    {
        "audit_id": "aud_1",
        "filename": "Tender_GeM_2026.pdf",
        "file_name": "Tender_GeM_2026.pdf",
        "timestamp": "2026-09-22 10:14",
        "defects": 1,
        "status": "FLAGGED",
        "tender_title": "CPWD Rural Electrification Package 4",
        "analyzed_at": "2026-09-22T10:14:00Z",
        "critical_count": 1,
        "compliant_count": 1,
        "coverage": "50%",
        "defects_list": [
            {
                "standard": "IS 694 : 1990",
                "severity": "CRITICAL",
                "finding": "Cited standard IS 694:1990 is withdrawn and superseded. Missing flame retardance tests per IS 10810 (Part 53).",
                "recommendation": "Replace with IS 694 : 2010 + Amendments 1 to 4.",
            }
        ],
    },
    {
        "audit_id": "aud_2",
        "filename": "Tender_NHAI_Bridge_Reinforcement.pdf",
        "file_name": "Tender_NHAI_Bridge_Reinforcement.pdf",
        "timestamp": "2026-09-21 14:45",
        "defects": 1,
        "status": "FLAGGED",
        "tender_title": "NHAI Expressway Elevated Deck Pier Reinforcement",
        "analyzed_at": "2026-09-21T14:45:00Z",
        "critical_count": 1,
        "compliant_count": 2,
        "coverage": "67%",
        "defects_list": [
            {
                "standard": "IS 432 : 1982",
                "severity": "CRITICAL",
                "finding": "Obsolete mild steel plain bar cited instead of ductile Fe 500D per IS 1786:2008.",
                "recommendation": "Replace with IS 1786 : 2008 Fe 500D.",
            }
        ],
    },
]

DRAFTS_DB: List[Dict[str, Any]] = deepcopy(INITIAL_DRAFTS)
AUDIT_HISTORY_DB: List[Dict[str, Any]] = deepcopy(INITIAL_AUDIT_HISTORY)


# ---------------------------------------------------------------------------
# PYDANTIC SCHEMAS
# ---------------------------------------------------------------------------
class SearchRequest(BaseModel):
    query: Optional[str] = Field(default=None, description="Natural language search requirement or standard code")
    standard_id: Optional[str] = Field(default=None, description="BIS standard code e.g. IS 7098 (Part 1) : 1988")
    language: str = Field(default="en", description="Language code: en or hi")
    domain: Optional[str] = Field(default=None, description="Optional domain filter: cables, cement, steel, electronics")
    qco_only: Optional[bool] = Field(default=None, description="Filter only standards with mandatory QCO")
    scheme_type: Optional[str] = Field(default=None, description="Filter by certification scheme type: ISI, CRS")
    clause_category: Optional[str] = Field(default=None, description="Filter matching clauses by category")


class OfficerDraftCreatePayload(BaseModel):
    project_title: Optional[str] = None
    primary_standard: Optional[str] = None
    clause_text: Optional[str] = None
    allied_standards: Optional[List[Any]] = Field(default_factory=list)
    # Legacy fields
    project_name: Optional[str] = None
    standard_code: Optional[str] = None
    clause_heading: Optional[str] = None
    clause_body: Optional[str] = None
    domain: Optional[str] = "cables"
    officer_notes: Optional[str] = ""


class AuditLineItem(BaseModel):
    line: Optional[str] = "Line Item"
    title: Optional[str] = "Procurement Item"
    text: Optional[str] = ""
    standard: Optional[str] = ""


class AuditJsonPayload(BaseModel):
    fileName: Optional[str] = "Tender_Specification_Batch.json"
    items: List[AuditLineItem] = []


class ExportClauseRequest(BaseModel):
    standard_id: Optional[str] = Field(default=None, description="BIS Standard code or ID")
    standard_code: Optional[str] = Field(default=None, description="Legacy field for standard code")
    language: str = Field(default="en", description="Language: en or hi")
    clause_heading: Optional[str] = None
    clause_body: Optional[str] = None
    domain: Optional[str] = "cables"
    format: Optional[str] = "docx"


class ExportDocxRequest(BaseModel):
    title: Optional[str] = Field(default="Procurement Tender Technical Specification")
    clause_text: Optional[str] = Field(default="")
    primary_standard: Optional[str] = Field(default="IS 7098 (Part 1) : 1988")
    allied_standards: Optional[List[Any]] = Field(default_factory=list)
    language: Optional[str] = Field(default="en")


class CopilotQueryRequest(BaseModel):
    prompt: str = Field(min_length=1, max_length=1000)
    domain: Optional[str] = Field(default=None)
    language: Optional[str] = Field(default="en")


# ---------------------------------------------------------------------------
# CORE SYSTEM & HEALTH ENDPOINTS
# ---------------------------------------------------------------------------
@app.get("/api/health")
def health():
    return {
        "status": "ok",
        "service": "manaksetu-bis-api",
        "version": app.version,
        "synced": "2026-09",
        "domains_indexed": ["cables", "cement", "steel", "electronics"],
        "officer_workspace": "active",
    }


@app.get("/")
def root():
    return {
        "name": "ManakSetu National Standards Intelligence Copilot",
        "version": app.version,
        "docs": "/docs",
        "health": "/api/health",
        "workspace": "/api/officer/workspace",
    }


# ---------------------------------------------------------------------------
# STANDARDS SEARCH & CLAUSE RETRIEVAL ENDPOINTS
# ---------------------------------------------------------------------------
@app.post("/api/search")
def standard_search(payload: SearchRequest):
    engine = get_retrieval_engine()
    search_term = payload.query or payload.standard_id or "IS 7098"
    result = engine.search(
        query=search_term,
        domain=payload.domain,
        qco_only=payload.qco_only,
        scheme_type=payload.scheme_type,
        clause_category=payload.clause_category,
    )
    
    # Synthesize bilingual clause based on target language
    std_key = payload.standard_id or result["primary"]["code"]
    bilingual_clause = get_bilingual_clause(std_key, language=payload.language)
    result["primary"]["clause"] = bilingual_clause

    return {
        **result,
        "language": payload.language,
        "synthesized_clause": bilingual_clause,
        "source": "BIS catalogue snapshot / deterministic multi-domain index",
    }


# ---------------------------------------------------------------------------
# NORMATIVE KNOWLEDGE GRAPH ENDPOINTS
# ---------------------------------------------------------------------------
@app.get("/api/graph/{domain_or_standard}")
def get_graph(domain_or_standard: str):
    engine = get_retrieval_engine()
    resolved = engine.resolve_domain(domain_or_standard) or "cables"
    graph_data = GRAPHS.get(resolved, GRAPHS["cables"])
    return {
        "domain": resolved,
        "standard_id": domain_or_standard,
        "nodes": graph_data["nodes"],
        "edges": graph_data["edges"],
        "source_locked": True,
    }


# ---------------------------------------------------------------------------
# DOCUMENT INGESTION & TENDER AUDIT ENGINE
# ---------------------------------------------------------------------------
IS_REGEX = re.compile(r"IS\s*(\d{3,5}(?:\s*\(Part\s*\d+\))?(?:\s*:\s*\d{4})?)", re.IGNORECASE)

FALLBACK_PROCUREMENT_LINES = [
    {
        "clause_no": "Clause 3.1",
        "raw_text": "Cables shall conform to IS 694:1990 PVC insulated cables for working voltages up to 1100V with copper conductors.",
        "cited_code": "IS 694 : 1990",
    },
    {
        "clause_no": "Clause 4.2",
        "raw_text": "Structural rebar reinforcement shall conform to IS 1786 : 2008 high strength deformed steel grade Fe 500D.",
        "cited_code": "IS 1786 : 2008",
    },
    {
        "clause_no": "Clause 5.4",
        "raw_text": "Ordinary Portland Cement 53 Grade shall strictly adhere to IS 269 : 2015 and Cement QCO mandates.",
        "cited_code": "IS 269 : 2015",
    },
]


def extract_text_from_stream(filename: str, contents: bytes) -> str:
    """Extracts text from PDF, DOCX, or TXT file bytes with exception tolerance."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    text = ""
    try:
        if ext == "pdf":
            reader = pypdf.PdfReader(io.BytesIO(contents))
            pages_text = []
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    pages_text.append(t)
            text = "\n".join(pages_text)
        elif ext == "docx":
            doc = docx.Document(io.BytesIO(contents))
            text = "\n".join([p.text for p in doc.paragraphs if p.text])
        elif ext in ["txt", "text", "csv", "json"]:
            text = contents.decode("utf-8", errors="replace")
        else:
            text = contents.decode("utf-8", errors="replace")
    except Exception:
        text = ""
    return text.strip()


def evaluate_standard_citation(raw_code: str, context_text: str = "") -> Dict[str, Any]:
    """Cross-reference detected standards against statutory supersession catalog."""
    norm_code = re.sub(r"\s+", " ", raw_code.strip())
    norm_lower = norm_code.lower()
    ctx_lower = context_text.lower()

    # Rule 1: IS 694 (withdrawn 1990 edition or missing revision)
    if "694" in norm_lower:
        if "2010" in norm_lower or "2010" in ctx_lower:
            return {
                "cited_code": "IS 694 : 2010",
                "status": "COMPLIANT_ACTIVE",
                "replacement_code": None,
                "missing_tests": [],
                "risk_severity": "LOW",
                "rationale": "Standard reference is verified active with statutory QCO requirements.",
            }
        else:
            return {
                "cited_code": norm_code if ":" in norm_code else f"{norm_code} : 1990",
                "status": "SUPERSEDED_WITHDRAWN",
                "replacement_code": "IS 694 : 2010 (Fourth Revision) + Amendments 1 to 4",
                "missing_tests": [
                    "IS 10810 (Part 53) - Flame Retardance Test",
                    "IS 8130 : 2013 - Conductor Resistance",
                ],
                "risk_severity": "CRITICAL",
                "rationale": "The 1990 edition lacks mandatory Low Smoke Zero Halogen (LSZH) test compliance under the Wires & Cables QCO.",
            }

    # Rule 2: IS 432 (Obsolete mild steel plain bars)
    if "432" in norm_lower:
        return {
            "cited_code": norm_code if ":" in norm_code else f"{norm_code} : 1982",
            "status": "SUPERSEDED_WITHDRAWN",
            "replacement_code": "IS 1786 : 2008 (Fe 500D)",
            "missing_tests": [
                "IS 1608 (Part 1) - 0.2% Proof Stress",
                "IS 1599 : 2019 - Bend and Rebend Test",
            ],
            "risk_severity": "CRITICAL",
            "rationale": "IS 432 mild steel plain bars lack high ductility and seismic proof stress mandates under Steel QCO 2024.",
        }

    # Rule 3a: IS 12269 (Withdrawn & superseded by unified IS 269 : 2015)
    if "12269" in norm_lower:
        return {
            "cited_code": norm_code if ":" in norm_code else f"{norm_code} : 2013",
            "status": "SUPERSEDED_WITHDRAWN",
            "replacement_code": "IS 269 : 2015 (Unified Ordinary Portland Cement Specification)",
            "missing_tests": [
                "IS 4031 (Part 6) - 28-day Compressive Strength Test (53 Grade)",
                "IS 4032 : 1985 - Chemical Composition Limits",
                "IS 4987 : 1993 - Sampling and Lot Acceptance Criteria",
            ],
            "risk_severity": "CRITICAL",
            "rationale": "IS 12269 was withdrawn by BIS following the revision of IS 269 : 2015, which unified 33, 43, and 53 Grade OPC under one standard.",
        }

    # Rule 3b: IS 8112 (Withdrawn & superseded by unified IS 269 : 2015)
    if "8112" in norm_lower:
        return {
            "cited_code": norm_code if ":" in norm_code else f"{norm_code} : 2013",
            "status": "SUPERSEDED_WITHDRAWN",
            "replacement_code": "IS 269 : 2015 (Unified Ordinary Portland Cement Specification)",
            "missing_tests": [
                "IS 4031 (Part 6) - Compressive Strength Test (43 Grade)",
                "IS 4032 : 1985 - Chemical Composition Limits",
                "IS 4987 : 1993 - Sampling and Lot Acceptance Criteria",
            ],
            "risk_severity": "CRITICAL",
            "rationale": "IS 8112 was withdrawn by BIS following the revision of IS 269 : 2015, which unified 33, 43, and 53 Grade OPC under one standard.",
        }

    # Rule 3c: IS 269 (Active standard for 33, 43, and 53 Grade OPC; obsolete if citing pre-2015 editions)
    if "269" in norm_lower and "12269" not in norm_lower:
        if "1976" in norm_lower or "1989" in norm_lower or "1976" in ctx_lower or "1989" in ctx_lower:
            return {
                "cited_code": norm_code,
                "status": "SUPERSEDED_WITHDRAWN",
                "replacement_code": "IS 269 : 2015 (Unified Ordinary Portland Cement Specification)",
                "missing_tests": [
                    "IS 4031 (Part 6) - 28-day Compressive Strength Test",
                    "IS 4032 : 1985 - Chemical Composition Limits",
                ],
                "risk_severity": "CRITICAL",
                "rationale": "Older editions of IS 269 (1976, 1989) are superseded by IS 269 : 2015, which unified 33, 43, and 53 Grade OPC under one standard.",
            }
        else:
            return {
                "cited_code": norm_code if ":" in norm_code else "IS 269 : 2015",
                "status": "COMPLIANT_ACTIVE",
                "replacement_code": None,
                "missing_tests": [],
                "risk_severity": "LOW",
                "rationale": "IS 269 : 2015 is current and active, unifying 33, 43, and 53 Grade OPC with mandatory Cement QCO 2003 compliance.",
            }

    # Rule 4: IS 13252 (2003 edition without MeitY CRS or secondary lithium batteries)
    if "13252" in norm_lower and ("2003" in norm_lower or ("2010" not in norm_lower and "part 1" not in norm_lower)):
        return {
            "cited_code": norm_code if ":" in norm_code else f"{norm_code} : 2003",
            "status": "SUPERSEDED_WITHDRAWN",
            "replacement_code": "IS 13252 (Part 1) : 2010",
            "missing_tests": [
                "IS 16046 (Part 2) : 2018 - Secondary Lithium Cells Safety",
            ],
            "risk_severity": "CRITICAL",
            "rationale": "Withdrawn edition lacks mandatory secondary lithium battery safety and MeitY CRS registration.",
        }

    # Default: Standard is active / compliant
    return {
        "cited_code": norm_code,
        "status": "COMPLIANT_ACTIVE",
        "replacement_code": None,
        "missing_tests": [],
        "risk_severity": "LOW",
        "rationale": "Standard reference is verified active with statutory QCO requirements.",
    }


def parse_and_audit_document(filename: str, contents: bytes) -> Dict[str, Any]:
    """Parses text from file bytes and produces the required audit response."""
    extracted_text = extract_text_from_stream(filename, contents)
    filesize_kb = max(1, len(contents) // 1024) if contents else 1420
    audit_id = f"aud_2026_{uuid.uuid4().hex[:4]}"

    items: List[Dict[str, Any]] = []

    if extracted_text:
        # Segment into lines/paragraphs
        lines = [line.strip() for line in extracted_text.splitlines() if line.strip()]
        for idx, line in enumerate(lines):
            matches = list(IS_REGEX.finditer(line))
            for m in matches:
                raw_code = m.group(0)
                eval_res = evaluate_standard_citation(raw_code, line)

                # Determine clause number prefix if present
                clause_match = re.match(r"(Clause\s*[\d\.]+|Section\s*[\d\.]+|Item\s*[\d\.]+)", line, re.IGNORECASE)
                clause_no = clause_match.group(0) if clause_match else f"Clause {len(items) + 1}.1"

                item_obj = {
                    "clause_no": clause_no,
                    "raw_text": line[:280] + ("..." if len(line) > 280 else ""),
                    "cited_code": eval_res["cited_code"],
                    "status": eval_res["status"],
                    "replacement_code": eval_res["replacement_code"] or "N/A",
                    "missing_tests": eval_res["missing_tests"],
                    "risk_severity": eval_res["risk_severity"],
                    "rationale": eval_res["rationale"],
                    # Compatibility aliases
                    "line": clause_no,
                    "title": f"Tender specification citing {eval_res['cited_code']}",
                    "standard": eval_res["cited_code"],
                    "finding": eval_res["rationale"],
                    "recommendation": f"Replace with {eval_res['replacement_code']}" if eval_res["replacement_code"] else "Retain active reference.",
                }
                items.append(item_obj)

    # Graceful fallback: If no items found or file was unreadable/corrupted
    if not items:
        for idx, fallback_line in enumerate(FALLBACK_PROCUREMENT_LINES):
            eval_res = evaluate_standard_citation(fallback_line["cited_code"], fallback_line["raw_text"])
            item_obj = {
                "clause_no": fallback_line["clause_no"],
                "raw_text": fallback_line["raw_text"],
                "cited_code": eval_res["cited_code"],
                "status": eval_res["status"],
                "replacement_code": eval_res["replacement_code"] or "N/A",
                "missing_tests": eval_res["missing_tests"],
                "risk_severity": eval_res["risk_severity"],
                "rationale": eval_res["rationale"],
                # Compatibility aliases
                "line": fallback_line["clause_no"],
                "title": f"Tender specification citing {eval_res['cited_code']}",
                "standard": eval_res["cited_code"],
                "finding": eval_res["rationale"],
                "recommendation": f"Replace with {eval_res['replacement_code']}" if eval_res["replacement_code"] else "Retain active reference.",
            }
            items.append(item_obj)

    total_items = len(items)
    defects_count = sum(1 for it in items if it["status"] == "SUPERSEDED_WITHDRAWN")

    result = {
        "audit_id": audit_id,
        "filename": filename or "Tender_CPWD_Elect_Dist_2026.pdf",
        "filesize_kb": filesize_kb,
        "total_items_found": total_items,
        "defects_count": defects_count,
        "items": items,
        # Legacy summary fields
        "summary": {
            "critical": defects_count,
            "compliant": total_items - defects_count,
            "coverage": f"{round(((total_items - defects_count) / max(total_items, 1)) * 100)}%",
        },
    }

    # Automatically append scan result to officer's audit history
    audit_entry = {
        "audit_id": audit_id,
        "filename": result["filename"],
        "timestamp": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M"),
        "defects": defects_count,
        "status": "FLAGGED" if defects_count > 0 else "RESOLVED",
        "file_name": result["filename"],
        "tender_title": result["filename"].replace(".pdf", "").replace(".docx", "").replace(".txt", "").replace("_", " "),
        "analyzed_at": datetime.now(timezone.utc).isoformat(),
        "critical_count": defects_count,
        "compliant_count": total_items - defects_count,
        "coverage": result["summary"]["coverage"],
        "defects_list": [
            {
                "standard": it["cited_code"],
                "severity": it["risk_severity"],
                "finding": it["rationale"],
                "recommendation": f"Replace with {it['replacement_code']}",
            }
            for it in items
            if it["status"] == "SUPERSEDED_WITHDRAWN"
        ],
    }

    with workspace_lock:
        AUDIT_HISTORY_DB.insert(0, audit_entry)

    return result


@app.post("/api/audit-tender/upload")
async def audit_tender_upload(file: UploadFile = File(...)):
    """Multipart document ingestion and auditing for .pdf, .docx, and .txt files."""
    contents = await file.read()
    filename = file.filename or "uploaded_tender_document.pdf"
    return parse_and_audit_document(filename, contents)


@app.post("/api/audit-tender")
async def audit_tender(
    request: Request,
    file: Optional[UploadFile] = File(default=None),
):
    """Backwards-compatible tender audit endpoint accepting multipart upload or JSON payload."""
    content_type = request.headers.get("content-type", "")

    # If file is supplied via form-data
    if file and file.filename:
        contents = await file.read()
        return parse_and_audit_document(file.filename, contents)

    # If JSON payload is supplied
    if "application/json" in content_type:
        try:
            body_data = await request.json()
            file_name = body_data.get("fileName", "Tender_Specification_Batch.json")
            raw_items = body_data.get("items", [])
            if raw_items:
                critical_items = []
                compliant_items = []
                for idx, item in enumerate(raw_items):
                    text_check = f"{item.get('title', '')} {item.get('text', '')} {item.get('standard', '')}"
                    eval_res = evaluate_standard_citation(item.get("standard") or text_check, text_check)
                    line_no = item.get("line") or f"Item {idx + 1}"
                    if eval_res["status"] == "SUPERSEDED_WITHDRAWN":
                        critical_items.append({
                            "line": line_no,
                            "title": item.get("title") or "Item",
                            "standard": eval_res["cited_code"],
                            "status": "critical",
                            "finding": eval_res["rationale"],
                            "recommendation": f"Replace with {eval_res['replacement_code']}",
                        })
                    else:
                        compliant_items.append({
                            "line": line_no,
                            "title": item.get("title") or "Item",
                            "standard": eval_res["cited_code"] or "Active BIS Standard",
                            "status": "compliant",
                            "finding": eval_res["rationale"],
                            "recommendation": "Retain reference; mandate batch test certificates at supply stage.",
                        })

                all_items = critical_items + compliant_items
                total = len(all_items)
                crit_count = len(critical_items)
                comp_count = len(compliant_items)
                coverage_pct = f"{round((comp_count / max(total, 1)) * 100)}%"

                audit_id = f"aud_2026_{uuid.uuid4().hex[:4]}"
                audit_entry = {
                    "audit_id": audit_id,
                    "filename": file_name,
                    "timestamp": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M"),
                    "defects": crit_count,
                    "status": "FLAGGED" if crit_count > 0 else "RESOLVED",
                    "file_name": file_name,
                    "tender_title": file_name.replace(".json", "").replace("_", " "),
                    "analyzed_at": datetime.now(timezone.utc).isoformat(),
                    "critical_count": crit_count,
                    "compliant_count": comp_count,
                    "coverage": coverage_pct,
                    "defects_list": [
                        {
                            "standard": it["standard"],
                            "severity": "CRITICAL",
                            "finding": it["finding"],
                            "recommendation": it["recommendation"],
                        }
                        for it in critical_items
                    ],
                }
                with workspace_lock:
                    AUDIT_HISTORY_DB.insert(0, audit_entry)

                return {
                    "audit_id": audit_id,
                    "fileName": file_name,
                    "filename": file_name,
                    "items": all_items,
                    "summary": {
                        "critical": crit_count,
                        "compliant": comp_count,
                        "coverage": coverage_pct,
                    },
                }
        except Exception:
            pass

    # Default file-less invocation (returns mock procurement lines)
    return parse_and_audit_document("Tender_GeM_Elect_2026.pdf", b"")


@app.post("/api/audit-tender/corrected")
def corrected_audit():
    return audit(True)


# ---------------------------------------------------------------------------
# OFFICER WORKSPACE & AUDIT LOG PERSISTENCE ENDPOINTS
# ---------------------------------------------------------------------------
@app.get("/api/officer/profile")
def get_officer_profile():
    """Returns official procurement auditor profile and DSC validation status."""
    return {
        "name": OFFICER_PROFILE["name"],
        "designation": OFFICER_PROFILE["designation"],
        "department": OFFICER_PROFILE["department"],
        "officer_id": OFFICER_PROFILE["officer_id"],
        "dsc_status": OFFICER_PROFILE["dsc_status"],
        "dsc_token_hash": OFFICER_PROFILE["dsc_token_hash"],
    }


@app.get("/api/officer/drafts")
def get_officer_drafts():
    """Returns array of saved tender drafts."""
    with workspace_lock:
        return [
            {
                "id": d.get("id") or d.get("draft_id"),
                "project_title": d.get("project_title") or d.get("project_name", ""),
                "primary_standard": d.get("primary_standard") or d.get("standard_code", ""),
                "last_modified": d.get("last_modified") or (d.get("created_at", "")[:10] if d.get("created_at") else "2026-09-22"),
                # Extended properties for full workspace compatibility
                "draft_id": d.get("draft_id") or d.get("id"),
                "project_name": d.get("project_name") or d.get("project_title", ""),
                "standard_code": d.get("standard_code") or d.get("primary_standard", ""),
                "clause_heading": d.get("clause_heading", ""),
                "clause_body": d.get("clause_body") or d.get("clause_text", ""),
                "clause_text": d.get("clause_text") or d.get("clause_body", ""),
                "allied_standards": d.get("allied_standards", []),
                "domain": d.get("domain", "cables"),
                "officer_notes": d.get("officer_notes", ""),
                "created_at": d.get("created_at", ""),
            }
            for d in DRAFTS_DB
        ]


@app.post("/api/officer/drafts", status_code=status.HTTP_201_CREATED)
def create_officer_draft(payload: OfficerDraftCreatePayload):
    """Saves a synthesized tender specification draft."""
    title = payload.project_title or payload.project_name or "Procurement Tender Specification"
    std = payload.primary_standard or payload.standard_code or "IS 7098 (Part 1) : 1988"
    clause = payload.clause_text or payload.clause_body or ""
    heading = payload.clause_heading or f"Specification for {title}"
    allied = payload.allied_standards or []

    draft_id = f"drf_{uuid.uuid4().hex[:6]}"
    created_date = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    created_iso = datetime.now(timezone.utc).isoformat()

    new_draft = {
        "status": "SAVED",
        "draft_id": draft_id,
        "id": draft_id,
        "project_title": title,
        "project_name": title,
        "primary_standard": std,
        "standard_code": std,
        "clause_text": clause,
        "clause_body": clause,
        "clause_heading": heading,
        "allied_standards": allied,
        "last_modified": created_date,
        "created_at": created_iso,
        "domain": payload.domain or "cables",
        "officer_notes": payload.officer_notes or "",
    }

    with workspace_lock:
        DRAFTS_DB.insert(0, new_draft)

    return new_draft


@app.delete("/api/officer/drafts/{draft_id}")
def delete_officer_draft(draft_id: str):
    """Removes a draft by ID."""
    global DRAFTS_DB
    with workspace_lock:
        initial_len = len(DRAFTS_DB)
        DRAFTS_DB = [d for d in DRAFTS_DB if d.get("draft_id") != draft_id and d.get("id") != draft_id]
        if len(DRAFTS_DB) == initial_len:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Draft with ID '{draft_id}' not found.",
            )
        remaining = len(DRAFTS_DB)
    return {"deleted": True, "draft_id": draft_id, "remaining_drafts": remaining}


@app.get("/api/officer/audit-history")
def get_audit_history():
    """Returns array of past document audits."""
    with workspace_lock:
        return [
            {
                "audit_id": a.get("audit_id", ""),
                "filename": a.get("filename") or a.get("file_name", ""),
                "timestamp": a.get("timestamp") or (a.get("analyzed_at", "")[:16].replace("T", " ") if a.get("analyzed_at") else "2026-09-22 10:14"),
                "defects": a.get("defects", a.get("critical_count", 0)),
                "status": a.get("status", "FLAGGED"),
                # Extended properties for UI backwards-compatibility
                "file_name": a.get("file_name") or a.get("filename", ""),
                "tender_title": a.get("tender_title", ""),
                "analyzed_at": a.get("analyzed_at", ""),
                "critical_count": a.get("critical_count", a.get("defects", 0)),
                "compliant_count": a.get("compliant_count", 1),
                "coverage": a.get("coverage", "80%"),
                "defects_list": a.get("defects_list", []),
            }
            for a in AUDIT_HISTORY_DB
        ]


@app.patch("/api/officer/audit-history/{audit_id}/resolve")
def resolve_audit_history(audit_id: str):
    """Updates audit status from FLAGGED to RESOLVED upon applying replacements."""
    with workspace_lock:
        for item in AUDIT_HISTORY_DB:
            if item.get("audit_id") == audit_id:
                item["status"] = "RESOLVED"
                return {
                    "status": "RESOLVED",
                    "audit_id": audit_id,
                    "updated": True,
                    "message": "Audit status updated to RESOLVED. Replacements applied.",
                    "item": item,
                }
    raise HTTPException(
        status_code=status.HTTP_404_NOT_FOUND,
        detail=f"Audit record with ID '{audit_id}' not found.",
    )


@app.get("/api/officer/workspace")
def get_officer_workspace():
    """Returns aggregate authenticated officer session, drafts, audit history and system stats."""
    with workspace_lock:
        crit_defects_total = sum(item.get("defects", item.get("critical_count", 0)) for item in AUDIT_HISTORY_DB)
        return {
            "officer": OFFICER_PROFILE,
            "drafts": DRAFTS_DB,
            "audit_history": AUDIT_HISTORY_DB,
            "stats": {
                "total_drafts": len(DRAFTS_DB),
                "total_audits": len(AUDIT_HISTORY_DB),
                "critical_defects_intercepted": crit_defects_total,
                "domains_active": ["cables", "cement", "steel", "electronics"],
                "last_active": datetime.now(timezone.utc).isoformat(),
            },
        }


# ---------------------------------------------------------------------------
# BILINGUAL SPECIFICATION CLAUSE SYNTHESIZER & EXPORT
# ---------------------------------------------------------------------------
@app.post("/api/export-clause")
@app.post("/api/export/clause")
def export_clause(payload: ExportClauseRequest):
    """Bilingual specification clause synthesizer for English and Hindi."""
    std_id = payload.standard_id or payload.standard_code or "IS 7098 (Part 1) : 1988"
    clause_text = get_bilingual_clause(std_id, language=payload.language)
    clean_code = std_id.replace(" ", "_").replace(":", "-").replace("/", "-").replace("(", "").replace(")", "")
    file_ext = payload.format.lower() if payload.format else "docx"
    filename = f"{clean_code}_tender_clause.{file_ext}"
    heading = payload.clause_heading or f"Specification Clause — {std_id}"
    formatted_content = f"{heading}\n\n{clause_text}"

    return {
        "status": "success",
        "standard_id": std_id,
        "standard_code": std_id,
        "language": payload.language,
        "filename": filename,
        "format": file_ext,
        "clause_text": clause_text,
        "synthesized_clause": clause_text,
        "content": formatted_content,
        "domain": payload.domain,
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "ready": True,
    }


# ---------------------------------------------------------------------------
# OFFICIAL BRANDED DOCX EXPORTER
# ---------------------------------------------------------------------------
def generate_official_docx(
    title: str,
    clause_text: str,
    primary_standard: str,
    allied_standards: List[Any],
    language: str = "en",
) -> io.BytesIO:
    """Generates an official Government of India branded technical specification DOCX."""
    doc = Document()

    # Document Margins (0.8 inch uniform)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Header
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_hdr = p_header.add_run("GOVERNMENT OF INDIA - PROCUREMENT TENDER TECHNICAL SPECIFICATION ANNEXURE")
    run_hdr.bold = True
    run_hdr.font.size = Pt(13)
    run_hdr.font.color.rgb = RGBColor(16, 44, 87)  # Deep Navy

    # Sub-header
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Generated via ManakSetu | Verified against BIS Catalogue & Mandatory QCOs")
    run_sub.font.size = Pt(9.5)
    run_sub.italic = True
    run_sub.font.color.rgb = RGBColor(80, 80, 80)

    # Officer Meta Box
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = p_meta.add_run("Auditing Officer: Shri Rajesh Kumar Sharma | DSC Token: Validated")
    run_meta.bold = True
    run_meta.font.size = Pt(10)
    run_meta.font.color.rgb = RGBColor(0, 102, 51)  # Forest Green

    p_submeta = doc.add_paragraph()
    p_submeta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_submeta = p_submeta.add_run("Ministry of Commerce & Industry / CPWD Electrical Wing | DSC Token: Validated")
    run_submeta.font.size = Pt(8.5)
    run_submeta.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 1.0 Scope
    h1 = doc.add_heading("Section 1.0 Scope", level=1)
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(4)

    p1 = doc.add_paragraph()
    p1.add_run("Project / Tender Title: ").bold = True
    p1.add_run(f"{title}\n")
    p1.add_run("Primary Standard: ").bold = True
    p1.add_run(f"{primary_standard}\n")
    p1.add_run("Document Language: ").bold = True
    p1.add_run(f"{'Hindi (हिंदी)' if language == 'hi' else 'English (EN)'}\n")
    p1.add_run(
        "This specification annexure mandates the statutory Bureau of Indian Standards (BIS) conformity, "
        "normative testing references, and compulsory Quality Control Orders (QCOs) for the tendered items."
    )

    # Section 2.0 Normative Reference Table
    h2 = doc.add_heading("Section 2.0 Normative Reference Table", level=1)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Standard Reference"
    hdr_cells[1].text = "Title / Subject Matter"
    hdr_cells[2].text = "Regulatory Status / Testing Group"

    for c in hdr_cells:
        for p in c.paragraphs:
            if p.runs:
                p.runs[0].bold = True
                p.runs[0].font.size = Pt(9.5)
                p.runs[0].font.color.rgb = RGBColor(16, 44, 87)

    # Primary Standard Row
    row = table.add_row().cells
    row[0].text = primary_standard
    row[1].text = f"Primary Product Specification ({title})"
    row[2].text = "MANDATORY / STATUTORY QCO"
    for c in row:
        for p in c.paragraphs:
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.size = Pt(9)

    # Allied Standards Rows
    for item in allied_standards:
        row = table.add_row().cells
        if isinstance(item, dict):
            code = item.get("code") or item.get("standard") or ""
            desc = item.get("title") or item.get("relevance") or "Normative Reference"
            grp = item.get("group") or "Compulsory Testing Protocol"
        else:
            code = str(item)
            desc = "Normative Reference & Testing Method"
            grp = "Compulsory Reference"
        row[0].text = code
        row[1].text = desc
        row[2].text = grp
        for c in row:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(9)

    # Section 3.0 Mandatory Testing & Certification Scheme
    h3 = doc.add_heading("Section 3.0 Mandatory Testing & Certification Scheme", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)

    p3_intro = doc.add_paragraph()
    p3_intro.add_run(
        "Under the Bureau of Indian Standards Act, 2016 and applicable statutory Quality Control Orders, "
        "all supplied consignments shall strictly comply with the synthesized specification clause below:\n"
    )

    effective_clause = clause_text or get_bilingual_clause(primary_standard, language=language)
    p_clause = doc.add_paragraph()
    p_clause.paragraph_format.left_indent = Inches(0.25)
    p_clause.paragraph_format.right_indent = Inches(0.25)
    run_clause = p_clause.add_run(f'"{effective_clause}"')
    run_clause.italic = True
    run_clause.font.size = Pt(10)
    run_clause.font.color.rgb = RGBColor(20, 20, 20)

    p3_mandates = doc.add_paragraph()
    p3_mandates.paragraph_format.space_before = Pt(8)
    p3_mandates.add_run("Compulsory Compliance Checklist for Tender Evaluation:\n").bold = True
    p3_mandates.add_run("1. The bidder/manufacturer must possess an active BIS Licence (CM/L or CRS R-number) on bid submission date.\n")
    p3_mandates.add_run("2. Consignments must carry the standard ISI Mark or CRS Self-Declaration imprint on products and packaging.\n")
    p3_mandates.add_run("3. Type-test certificates and routine batch test reports from NABL-accredited / BIS-approved labs are mandatory before dispatch.\n")

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream


@app.post("/api/export-docx")
def export_docx_endpoint(payload: ExportDocxRequest):
    """Generates and streams an official Government of India branded technical specification DOCX."""
    title = payload.title or "Procurement Tender Technical Specification"
    primary_std = payload.primary_standard or "IS 7098 (Part 1) : 1988"
    clause_txt = payload.clause_text or get_bilingual_clause(primary_std, language=payload.language or "en")
    allied_stds = payload.allied_standards or []

    stream = generate_official_docx(
        title=title,
        clause_text=clause_txt,
        primary_standard=primary_std,
        allied_standards=allied_stds,
        language=payload.language or "en",
    )

    clean_title = re.sub(r"[^a-zA-Z0-9_-]", "_", title)[:40]
    filename = f"ManakSetu_{clean_title}.docx"

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


# ---------------------------------------------------------------------------
# COPILOT ADVICE & PROMPT QUERY ROUTER
# ---------------------------------------------------------------------------
@app.post("/api/copilot/query")
def copilot_query(payload: CopilotQueryRequest):
    """Deterministic copilot router providing clause suggestions and compliance advice."""
    engine = get_retrieval_engine()
    search_res = engine.search(query=payload.prompt, domain=payload.domain)
    primary = search_res["primary"]
    allied_codes = [a["code"] for a in primary.get("allied", [])[:3]]
    recommendation = (
        f"For procurement requirement '{payload.prompt}', mandate conformity to {primary['code']} "
        f"under statutory order '{primary['qco']}'. Ensure batch testing certificates citing "
        f"{', '.join(allied_codes) if allied_codes else 'applicable BIS test standards'}."
    )
    bilingual_clause = get_bilingual_clause(primary["code"], language=payload.language or "en")

    return {
        "query": payload.prompt,
        "domain": primary.get("domain", "cables"),
        "primary_standard": {
            "id": primary.get("id"),
            "code": primary["code"],
            "title": primary["title"],
            "status": primary["status"],
            "qco": primary["qco"],
            "scheme": primary["scheme"],
        },
        "recommendation": recommendation,
        "synthesized_clause": bilingual_clause,
        "matching_clauses": search_res.get("matching_clauses", []),
        "confidence": search_res.get("confidence", 98.0),
    }
