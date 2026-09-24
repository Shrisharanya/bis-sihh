import io
import docx
import pypdf
import pytest
from fastapi.testclient import TestClient

from main import app, DRAFTS_DB, AUDIT_HISTORY_DB

client = TestClient(app)


def test_health_and_root():
    res = client.get("/api/health")
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "ok"
    assert "cables" in data["domains_indexed"]
    assert "cement" in data["domains_indexed"]
    assert "steel" in data["domains_indexed"]
    assert "electronics" in data["domains_indexed"]

    root_res = client.get("/")
    assert root_res.status_code == 200
    assert "ManakSetu" in root_res.json()["name"]


def test_officer_profile():
    res = client.get("/api/officer/profile")
    assert res.status_code == 200
    data = res.json()
    assert data["name"] == "Shri Rajesh Kumar Sharma"
    assert data["designation"] == "Chief Procurement Officer & Standards Auditor"
    assert data["department"] == "Ministry of Commerce & Industry / CPWD Electrical Wing"
    assert data["officer_id"] == "#GOV-PROC-8821"
    assert data["dsc_status"] == "Active (NIC-CA Validated)"
    assert data["dsc_token_hash"].startswith("SHA256:7F89B")


def test_officer_drafts_routes():
    # 1. Get drafts list
    res = client.get("/api/officer/drafts")
    assert res.status_code == 200
    drafts = res.json()
    assert isinstance(drafts, list)
    assert len(drafts) >= 3
    sample = drafts[0]
    assert "id" in sample
    assert "project_title" in sample
    assert "primary_standard" in sample
    assert "last_modified" in sample

    # 2. Create new draft via POST /api/officer/drafts
    payload = {
        "project_title": "NIT-204: 33kV Substation Cables",
        "primary_standard": "IS 7098 (Part 1) : 1988",
        "clause_text": "The supplied material shall strictly conform to IS 7098 (Part 1) : 1988...",
        "allied_standards": ["IS 8130 : 2013", "IS 10810 (Part 53)"],
    }
    create_res = client.post("/api/officer/drafts", json=payload)
    assert create_res.status_code == 201
    created = create_res.json()
    assert created["status"] == "SAVED"
    assert created["draft_id"].startswith("drf_")
    assert created["project_title"] == payload["project_title"]
    assert created["primary_standard"] == payload["primary_standard"]

    new_id = created["draft_id"]

    # 3. Verify in list
    drafts_after = client.get("/api/officer/drafts").json()
    assert any(d["id"] == new_id or d["draft_id"] == new_id for d in drafts_after)

    # 4. Delete draft
    del_res = client.delete(f"/api/officer/drafts/{new_id}")
    assert del_res.status_code == 200
    assert del_res.json()["deleted"] is True

    # 5. Confirm 404 for deleted draft
    del_res_404 = client.delete(f"/api/officer/drafts/{new_id}")
    assert del_res_404.status_code == 404


def test_officer_audit_history_and_resolve():
    # 1. Get audit history
    res = client.get("/api/officer/audit-history")
    assert res.status_code == 200
    history = res.json()
    assert isinstance(history, list)
    assert len(history) >= 2
    first = history[0]
    assert "audit_id" in first
    assert "filename" in first
    assert "timestamp" in first
    assert "defects" in first
    assert "status" in first

    target_id = first["audit_id"]

    # 2. Patch to resolve
    patch_res = client.patch(f"/api/officer/audit-history/{target_id}/resolve")
    assert patch_res.status_code == 200
    patch_data = patch_res.json()
    assert patch_data["status"] == "RESOLVED"
    assert patch_data["audit_id"] == target_id

    # Verify updated in history list
    history_after = client.get("/api/officer/audit-history").json()
    matched = [h for h in history_after if h["audit_id"] == target_id]
    assert len(matched) == 1
    assert matched[0]["status"] == "RESOLVED"

    # Test resolving non-existent audit ID
    res_404 = client.patch("/api/officer/audit-history/aud_nonexistent_9999/resolve")
    assert res_404.status_code == 404


def test_audit_tender_upload_pdf():
    # Construct a valid minimal PDF with IS 694 : 1990 citation
    pdf_raw = b"""%PDF-1.4
1 0 obj <</Type /Catalog /Pages 2 0 R>> endobj
2 0 obj <</Type /Pages /Kids [3 0 R] /Count 1>> endobj
3 0 obj <</Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R>> endobj
4 0 obj <</Type /Font /Subtype /Type1 /BaseFont /Helvetica>> endobj
5 0 obj <</Length 140>> stream
BT
/F1 12 Tf
72 712 Td
(Clause 3.1: Cables shall conform to IS 694:1990 PVC insulated cables for working voltages up to 1100V) Tj
ET
endstream
endobj
xref
0 6
0000000000 65535 f 
0000000010 00000 n 
0000000060 00000 n 
0000000117 00000 n 
0000000234 00000 n 
0000000305 00000 n 
trailer <</Size 6 /Root 1 0 R>>
startxref
496
%%EOF"""

    files = {"file": ("Tender_CPWD_Elect_Dist_2026.pdf", pdf_raw, "application/pdf")}
    res = client.post("/api/audit-tender/upload", files=files)
    assert res.status_code == 200
    data = res.json()

    # Verify JSON contract
    assert "audit_id" in data
    assert data["filename"] == "Tender_CPWD_Elect_Dist_2026.pdf"
    assert "filesize_kb" in data
    assert data["total_items_found"] >= 1
    assert data["defects_count"] >= 1
    assert len(data["items"]) >= 1

    item = data["items"][0]
    assert item["clause_no"].startswith("Clause")
    assert "IS 694" in item["cited_code"]
    assert item["status"] == "SUPERSEDED_WITHDRAWN"
    assert "IS 694 : 2010" in item["replacement_code"]
    assert item["risk_severity"] == "CRITICAL"
    assert any("10810" in t for t in item["missing_tests"])
    assert any("8130" in t for t in item["missing_tests"])
    assert "LSZH" in item["rationale"] or "Wires & Cables QCO" in item["rationale"]

    # Verify side-effect: audit automatically appended to audit history
    history = client.get("/api/officer/audit-history").json()
    assert any(h["audit_id"] == data["audit_id"] for h in history)


def test_audit_tender_upload_docx():
    # Construct a valid DOCX with IS 432:1982 citation
    doc = docx.Document()
    doc.add_paragraph("Clause 2.1: Reinforcement bars shall strictly conform to IS 432 : 1982 mild steel plain bars.")
    doc.add_paragraph("Clause 4.3: High strength steel shall comply with IS 1786 : 2008 Grade Fe 500D.")
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    files = {"file": ("Tender_NHAI_Structural.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = client.post("/api/audit-tender/upload", files=files)
    assert res.status_code == 200
    data = res.json()

    assert data["filename"] == "Tender_NHAI_Structural.docx"
    assert data["total_items_found"] >= 2
    assert data["defects_count"] >= 1

    superseded = [it for it in data["items"] if it["status"] == "SUPERSEDED_WITHDRAWN"]
    assert len(superseded) >= 1
    assert "IS 432" in superseded[0]["cited_code"]
    assert "IS 1786" in superseded[0]["replacement_code"]


def test_audit_tender_upload_txt():
    txt_content = (
        "Clause 1.0 General Building Construction\n"
        "Clause 5.1 Cement shall strictly conform to IS 269 : 1976 33 Grade Ordinary Portland Cement.\n"
        "Clause 5.2 High strength concrete shall use IS 12269 : 2013 53 Grade OPC."
    ).encode("utf-8")

    files = {"file": ("Tender_CPWD_Civil.txt", txt_content, "text/plain")}
    res = client.post("/api/audit-tender/upload", files=files)
    assert res.status_code == 200
    data = res.json()

    assert data["filename"] == "Tender_CPWD_Civil.txt"
    assert data["defects_count"] >= 1
    superseded = [it for it in data["items"] if it["status"] == "SUPERSEDED_WITHDRAWN"]
    assert "IS 269" in superseded[0]["cited_code"]
    assert "IS 12269" in superseded[0]["replacement_code"]


def test_audit_tender_upload_corrupted_fallback():
    # Send corrupted binary bytes that cannot be parsed as PDF
    corrupt_bytes = b"CORRUPTED_GARBAGE_BINARY_STREAM_%$#@!*&^"
    files = {"file": ("Corrupted_Tender.pdf", corrupt_bytes, "application/pdf")}
    res = client.post("/api/audit-tender/upload", files=files)
    assert res.status_code == 200
    data = res.json()

    # Graceful fallback returns mock procurement lines
    assert data["total_items_found"] == 3
    assert data["defects_count"] == 1
    assert data["items"][0]["status"] == "SUPERSEDED_WITHDRAWN"


def test_bilingual_clause_synthesizer_en_and_hi():
    # 1. English synthesis in POST /api/search
    res_en = client.post("/api/search", json={"standard_id": "IS 7098 (Part 1) : 1988", "language": "en"})
    assert res_en.status_code == 200
    data_en = res_en.json()
    assert data_en["language"] == "en"
    clause_en = data_en["synthesized_clause"]
    assert "The supplied material shall strictly conform to IS 7098 (Part 1) : 1988" in clause_en
    assert "IS 8130 : 2013" in clause_en
    assert "IS 10810 (Series)" in clause_en
    assert "Wires and Cables (Quality Control) Order, 2023" in clause_en

    # 2. Hindi synthesis in POST /api/search
    res_hi = client.post("/api/search", json={"standard_id": "IS 7098 (Part 1) : 1988", "language": "hi"})
    assert res_hi.status_code == 200
    data_hi = res_hi.json()
    assert data_hi["language"] == "hi"
    clause_hi = data_hi["synthesized_clause"]
    assert "आपूर्त सामग्री को नवीनतम संशोधनों सहित IS 7098 (Part 1) : 1988 के पूर्णतः अनुरूप होना अनिवार्य है।" in clause_hi
    assert "IS 8130 : 2013" in clause_hi
    assert "IS 10810 शृंखला" in clause_hi
    assert "तार एवं केबल (गुणवत्ता नियंत्रण) आदेश, 2023" in clause_hi

    # 3. Bilingual synthesis in POST /api/export-clause
    exp_en = client.post("/api/export-clause", json={"standard_id": "IS 7098 (Part 1) : 1988", "language": "en"})
    assert exp_en.status_code == 200
    assert "IS 7098 (Part 1) : 1988" in exp_en.json()["synthesized_clause"]

    exp_hi = client.post("/api/export-clause", json={"standard_id": "IS 7098 (Part 1) : 1988", "language": "hi"})
    assert exp_hi.status_code == 200
    assert "आपूर्त सामग्री" in exp_hi.json()["synthesized_clause"]


def test_official_docx_export():
    payload = {
        "title": "NIT-204: 33kV Substation Cables Package",
        "primary_standard": "IS 7098 (Part 1) : 1988",
        "clause_text": "The supplied material shall strictly conform to IS 7098 (Part 1) : 1988 with mandatory ISI mark.",
        "allied_standards": [
            {"code": "IS 8130 : 2013", "title": "Conductors for insulated cables", "group": "Conductor Specs"},
            {"code": "IS 5831 : 1984", "title": "PVC insulation and sheath", "group": "Sheathing"},
            {"code": "IS 10810 (Part 53)", "title": "Vertical flame propagation", "group": "Fire Safety"},
        ],
        "language": "en",
    }
    res = client.post("/api/export-docx", json=payload)
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in res.headers["content-type"]
    assert "attachment;" in res.headers["content-disposition"]
    assert ".docx" in res.headers["content-disposition"]

    # Read back and inspect the binary docx content
    doc = docx.Document(io.BytesIO(res.content))
    doc_text = "\n".join([p.text for p in doc.paragraphs])

    # Verify Government of India header
    assert "GOVERNMENT OF INDIA - PROCUREMENT TENDER TECHNICAL SPECIFICATION ANNEXURE" in doc_text
    # Verify Sub-header
    assert "Generated via ManakSetu | Verified against BIS Catalogue & Mandatory QCOs" in doc_text
    # Verify Officer Meta
    assert "Auditing Officer: Shri Rajesh Kumar Sharma | DSC Token: Validated" in doc_text
    # Verify Section 1.0 Scope
    assert "Section 1.0 Scope" in doc_text
    # Verify Section 2.0 Normative Reference Table
    assert "Section 2.0 Normative Reference Table" in doc_text
    # Verify Section 3.0 Mandatory Testing & Certification Scheme
    assert "Section 3.0 Mandatory Testing & Certification Scheme" in doc_text

    # Verify table rows
    assert len(doc.tables) >= 1
    table_text = "\n".join([cell.text for row in doc.tables[0].rows for cell in row.cells])
    assert "IS 7098 (Part 1) : 1988" in table_text
    assert "IS 8130 : 2013" in table_text
    assert "IS 10810 (Part 53)" in table_text


def test_search_multi_domain_standards():
    # Cables
    res = client.post("/api/search", json={"query": "3.5 core XLPE cable 1.1kV", "language": "en"})
    assert res.status_code == 200
    assert res.json()["primary"]["code"] == "IS 7098 (Part 1) : 1988"

    # Cement
    res = client.post("/api/search", json={"query": "Ordinary Portland Cement 53 grade", "domain": "cement"})
    assert res.status_code == 200
    assert res.json()["primary"]["code"] == "IS 12269 : 2013"

    # Steel
    res = client.post("/api/search", json={"query": "Fe 500D TMT high strength deformed steel rebar"})
    assert res.status_code == 200
    assert res.json()["primary"]["code"] == "IS 1786 : 2008"

    # Electronics
    res = client.post("/api/search", json={"query": "IT Equipment Safety MeitY CRS hardware"})
    assert res.status_code == 200
    assert res.json()["primary"]["code"] == "IS 13252 (Part 1) : 2010"


def test_knowledge_graph_endpoints():
    domains = ["cables", "cement", "steel", "electronics"]
    for dom in domains:
        res = client.get(f"/api/graph/{dom}")
        assert res.status_code == 200
        data = res.json()
        assert data["domain"] == dom
        assert len(data["nodes"]) >= 5
        assert len(data["edges"]) >= 4


def test_copilot_query_endpoint():
    payload = {
        "prompt": "high strength deformed steel bars for seismic RCC columns",
        "domain": "steel",
    }
    res = client.post("/api/copilot/query", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert data["domain"] == "steel"
    assert "IS 1786" in data["primary_standard"]["code"]
    assert len(data["synthesized_clause"]) > 20
